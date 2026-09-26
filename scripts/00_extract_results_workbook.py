#!/usr/bin/env python3
"""Extract the analysis results used in the paper from the results workbook into CSV files.

The workbook ("Buriganga metagenomics_2025-01-12.xlsx") holds the OTU table, sequencing
statistics, alpha and beta diversity results, taxonomic abundance tables and PICRUSt2
predictions. This script writes the parts used in the paper to ../data/.

The 'OUT no' sheet stores the OTU table twice (two identical blocks); only the first
block is read, so every OTU appears once.

Usage:
    python3 00_extract_results_workbook.py "Buriganga metagenomics_2025-01-12.xlsx"
"""

import sys
from pathlib import Path

import openpyxl
import pandas as pd

HERE = Path(__file__).resolve().parent
DATA = HERE.parent / "data"

SAMPLES = [
    "DeulS", "NbbS", "AbbS", "BsbS", "MdcS", "KrcS", "BsgS", "JjfS", "BbbS",
    "CmbS", "PgbS", "HnbS", "PaglaS", "FatulS",
    "DeulW", "NbbW", "AbbW", "BsbW", "MdcW", "KrcW", "BsgW", "JjfW", "BbbW",
    "CmbW", "PgbW", "HnbW", "PaglaW", "FatulW",
]

# Site number (as in Fig. 1), name, river, zone and coordinates (WGS 84).
SITES = {
    "Deul":  (1, "Deul", "Turag", "Upstream", 23.8661403, 90.3507461),
    "Nbb":   (2, "Nobaberbag", "Turag", "Upstream", 23.8132878, 90.3393873),
    "Abb":   (3, "Aminbazar Bridge", "Turag", "Upstream", 23.7845523, 90.3363125),
    "Bsb":   (4, "Bosila Bridge", "Buriganga", "Midstream", 23.7440959, 90.3464334),
    "Mdc":   (5, "Madhyer Char", "Buriganga", "Midstream", 23.7296287, 90.3575465),
    "Krc":   (6, "Kamrangirchar", "Buriganga", "Midstream", 23.7099437, 90.3649564),
    "Bsg":   (7, "Barishur Ghat", "Buriganga", "Midstream", 23.7059582, 90.3759116),
    "Jjf":   (8, "Jinjira Ferryghat", "Buriganga", "Midstream", 23.7092334, 90.3933855),
    "Bbb":   (9, "Babubazar Bridge", "Buriganga", "Midstream", 23.7084813, 90.4004957),
    "Hnb":   (10, "Nayatola (Hasnabad)", "Buriganga", "Downstream", 23.6966650, 90.4121380),
    "Cmb":   (11, "Char Mirerbagh", "Buriganga", "Downstream", 23.6863720, 90.4257970),
    "Pgb":   (12, "Postagola Bridge", "Buriganga", "Downstream", 23.6879863, 90.4242685),
    "Pagla": (13, "Pagla", "Buriganga", "Downstream", 23.6609830, 90.4557016),
    "Fatul": (14, "Fatullah", "Buriganga", "Downstream", 23.6411826, 90.4723568),
}

RANKS = ["Kingdom", "Phylum", "Class", "Order", "Family", "Genus", "Species"]


def parse_taxonomy(tax):
    """Split a SILVA string like 'k__Bacteria;p__...;g__X' into seven ranks."""
    ranks = [""] * 7
    if not tax or tax.strip() == "Unknown":
        ranks[0] = "Unassigned"
        return ranks
    for i, part in enumerate(tax.split(";")[:7]):
        ranks[i] = part.split("__", 1)[-1].strip()
    return ranks


def otu_table(wb):
    rows = list(wb["OUT no"].iter_rows(values_only=True))
    header_row = next(i for i, r in enumerate(rows) if r[0] == "#OTU_num" and r[1] == "DeulS")
    assert list(rows[header_row][1:29]) == SAMPLES, "unexpected sample order"
    records, seen = [], set()
    for r in rows[header_row + 1:]:
        otu = r[0]
        if otu is None or not str(otu).startswith("OTU_"):
            if records:          # the first OTU block has ended
                break
            continue
        if otu in seen:
            break
        seen.add(otu)
        records.append([otu] + [int(x) for x in r[1:29]] + parse_taxonomy(r[29]) + [r[29] or "Unknown"])
    df = pd.DataFrame(records, columns=["OTU_ID"] + SAMPLES + RANKS + ["Taxonomy_SILVA138"])
    return df


def metadata():
    rows = []
    for s in SAMPLES:
        code, hab = s[:-1], s[-1]
        num, name, river, zone, lat, lon = SITES[code]
        rows.append(dict(Sample_ID=s, Site_code=code, Site_number=num, Site_name=name,
                         River=river, Zone=zone, Habitat="Sediment" if hab == "S" else "Water",
                         Latitude=lat, Longitude=lon, Collection="January 2023",
                         BioProject="PRJNA1399629"))
    return pd.DataFrame(rows)


def sequencing_stats(wb):
    qc = pd.DataFrame(list(wb["QC"].iter_rows(values_only=True))[1:], columns=[
        "Sample_ID", "Raw_PE_reads", "Merged_reads", "Quality_filtered_reads", "Non_chimeric_reads",
        "Bases_nt", "Mean_length_nt", "Q20_pct", "Q30_pct", "GC_pct", "Effective_pct"])
    qc = qc[qc.Sample_ID.isin(SAMPLES)]
    tag = pd.DataFrame(list(wb["Tag_stat"].iter_rows(max_col=6, values_only=True))[1:], columns=[
        "Sample_ID", "Total_tags", "Taxon_tags", "Unclassified_tags", "Unique_tags", "OTUs"])
    tag = tag[tag.Sample_ID.isin(SAMPLES)]
    out = qc.merge(tag, on="Sample_ID")
    return out.set_index("Sample_ID").loc[SAMPLES].reset_index()


def kegg_levels(wb):
    """Per-sample PICRUSt2 KEGG category abundances (levels 1-3) from 'FunctionPredict'."""
    rows = list(wb["FunctionPredict"].iter_rows(values_only=True))
    out = {}
    starts = [i for i, r in enumerate(rows) if r[0] in ("Level_1", "Level_2", "Level_3")]
    for i in starts:
        level = rows[i][0]
        assert list(rows[i][1:29]) == SAMPLES
        rec = []
        for r in rows[i + 1:]:
            if r[0] is None:
                break
            rec.append([r[0]] + [float(x) for x in r[1:29]])
        out[level] = pd.DataFrame(rec, columns=["Category"] + SAMPLES)
    return out


def two_tables(rows, first_col, second_col, n_first, n_second):
    """Two tables are placed side by side on these sheets; return both as DataFrames."""
    hdr = rows[0]
    left = [r for r in rows[1:] if r[first_col]]
    right = [r for r in rows[1:] if r[second_col]]
    a = pd.DataFrame([[r[first_col]] + [float(x) for x in r[first_col + 1:first_col + 1 + n_first]] for r in left],
                     columns=["Taxon"] + list(hdr[first_col + 1:first_col + 1 + n_first]))
    b = pd.DataFrame([[r[second_col]] + [float(x) for x in r[second_col + 1:second_col + 1 + n_second]] for r in right],
                     columns=["Taxon"] + list(hdr[second_col + 1:second_col + 1 + n_second]))
    return a, b


def alpha(wb):
    rows = list(wb["AlphaD."].iter_rows(values_only=True))
    cols = ["observed_species", "shannon", "simpson", "chao1", "ACE", "goods_coverage", "PD_whole_tree"]
    per_sample = pd.DataFrame([r[:8] for r in rows[1:29]], columns=["Sample_ID"] + cols)
    start = rows[0].index("group")                     # the group-mean table sits to the right
    group_means = pd.DataFrame([r[start:start + 8] for r in rows[1:7]], columns=["Group"] + cols)
    wilcox = pd.DataFrame([dict(Index=rows[48][c].replace("_Wilcox", ""), Comparison=rows[49][c], p=rows[49][c + 1])
                           for c in (0, 3, 6, 9, 12)])
    kw = []
    for r in rows[52:67]:
        for c, index in ((12, "observed_species"), (18, "shannon"), (24, "simpson")):
            if r[c]:
                kw.append(dict(Index=index, Comparison=r[c], Difference=r[c + 1], p=r[c + 2],
                               Significance=(r[c + 3] or "").strip(), LCL=r[c + 4], UCL=r[c + 5]))
    return per_sample, group_means, wilcox, pd.DataFrame(kw)


def adonis(wb):
    """Bray-Curtis Adonis (PERMANOVA) results."""
    rows = list(wb["BetaD."].iter_rows(values_only=True))
    out = []
    for r in rows[1:17] + rows[19:20]:
        if not r[0]:
            continue
        r2, resid = [float(x) for x in str(r[5]).replace(")", "").split("(")]
        out.append(dict(Comparison=r[0], Df=r[1], F=float(r[4]), R2=r2, R2_residual=resid, p=float(r[6])))
    return pd.DataFrame(out)


def main(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    DATA.mkdir(exist_ok=True)
    otus = otu_table(wb)
    assert otus.OTU_ID.is_unique
    otus.to_csv(DATA / "otu_table.csv", index=False)
    metadata().to_csv(DATA / "sample_metadata.csv", index=False)
    sequencing_stats(wb).to_csv(DATA / "sequencing_stats.csv", index=False)
    for level, df in kegg_levels(wb).items():
        df.to_csv(DATA / f"picrust2_kegg_{level.lower()}.csv", index=False)
    for sheet, name in (("Gen_Abun.", "genus"), ("Phyl_Abun.", "phylum")):
        habitat, groups = two_tables(list(wb[sheet].iter_rows(values_only=True)), 0, 4, 2, 6)
        habitat.rename(columns={"Taxon": name.capitalize()}).to_csv(DATA / f"{name}_habitat_means.csv", index=False)
        groups.rename(columns={"Taxon": name.capitalize()}).to_csv(DATA / f"{name}_group_means.csv", index=False)
    per_sample, group_means, wilcox, kw = alpha(wb)
    per_sample.to_csv(DATA / "alpha_diversity.csv", index=False)
    group_means.to_csv(DATA / "alpha_group_means.csv", index=False)
    wilcox.to_csv(DATA / "alpha_wilcoxon_sediment_vs_water.csv", index=False)
    kw.to_csv(DATA / "alpha_kruskal_pairwise.csv", index=False)
    adonis(wb).to_csv(DATA / "adonis_bray_curtis.csv", index=False)
    print(f"OTUs: {len(otus)}; reads in OTU table: {otus[SAMPLES].to_numpy().sum():,}")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "Buriganga metagenomics_2025-01-12.xlsx")
