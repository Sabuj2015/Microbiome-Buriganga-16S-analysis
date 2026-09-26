#!/usr/bin/env python3
"""Additional file 1: Tables S1-S3 and machine-readable versions of Tables 1-3.

Table S1  relative abundance of key genera by zone and habitat (mean, SD, min, max)
Table S2  sediment vs water comparison of the 108 genera with mean relative abundance >= 0.1%
          (Welch's t-test as in Fig. 5B; Mann-Whitney U test with Benjamini-Hochberg correction)
Table S3  predicted KEGG level-2 and level-3 categories (data underlying Fig. 6)
Tables 1-3 as in the manuscript

Relative abundance of a genus in a sample = reads assigned to the genus / reads in the OTU table.
Means in Table S1 are the group means of Table 3 (genus_group_means.csv); SD, minimum and
maximum are calculated from the per-sample relative abundances.
"""

from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from scipy import stats

HERE = Path(__file__).resolve().parent
DATA = HERE.parent / "data"
OUT = HERE.parent / "results"
OUT.mkdir(exist_ok=True)

GROUPS = {("Upstream", "Water"): "UpstrmW", ("Midstream", "Water"): "MidstrmW", ("Downstream", "Water"): "DownstmW",
          ("Upstream", "Sediment"): "UpstrmS", ("Midstream", "Sediment"): "MidstrmS", ("Downstream", "Sediment"): "DownstmS"}
S1_GENERA = ["Ralstonia", "Arcobacter", "Acinetobacter", "Cloacibacterium", "Aeromonas", "Bacillus", "Thiobacillus",
             "Thauera", "Sulfuricurvum", "Enterobacter", "Vibrio", "C39"]
T3 = [("Biogeochemical cycling", "Bacillus"), ("Biogeochemical cycling", "Thiobacillus"),
      ("Biogeochemical cycling", "Sulfuricurvum"), ("Biogeochemical cycling", "Thauera"),
      ("Biogeochemical cycling", "Enterobacter"), ("Potential pathogen", "Ralstonia"),
      ("Potential pathogen", "Acinetobacter"), ("Potential pathogen", "Arcobacter"),
      ("Potential pathogen", "Cloacibacterium"), ("Potential pathogen", "Aeromonas"),
      ("Redox / other", "Trichlorobacter"), ("Unclassified", "C39")]


def bh(p):
    p = np.asarray(p, float)
    order = np.argsort(p)
    q = p[order] * len(p) / np.arange(1, len(p) + 1)
    q = np.minimum.accumulate(q[::-1])[::-1]
    out = np.empty(len(p))
    out[order] = np.minimum(q, 1)
    return out


def genus_relative_abundance():
    otu = pd.read_csv(DATA / "otu_table.csv")
    meta = pd.read_csv(DATA / "sample_metadata.csv").set_index("Sample_ID")
    samples = list(meta.index)
    counts = otu[samples]
    genus = otu.Genus.fillna("").replace("", "Unassigned")
    rel = counts.groupby(genus).sum() / counts.sum()
    return rel * 100, meta


def table_s1(rel, meta):
    means = pd.read_csv(DATA / "genus_group_means.csv").set_index("Genus") * 100
    rows = []
    for g in S1_GENERA:
        for (zone, hab), col in GROUPS.items():
            ids = meta.index[(meta.Zone == zone) & (meta.Habitat == hab)]
            v = rel.loc[g, ids]
            m = means.loc[g, col]
            rows.append([g, f"{zone} {hab}", round(m, 2), round(v.std(ddof=1), 2), round(v.min(), 2),
                         round(v.max(), 2), f"{m:.2f} ± {v.std(ddof=1):.2f}"])
    return pd.DataFrame(rows, columns=["Genus", "Zone", "Mean (%)", "SD (%)", "Min (%)", "Max (%)", "Formatted"])


def table_s2(rel, meta):
    sed = meta.index[meta.Habitat == "Sediment"]
    wat = meta.index[meta.Habitat == "Water"]
    keep = rel.drop(index="Unassigned").loc[lambda d: d.mean(axis=1) >= 0.1]
    rows = []
    for g, v in keep.iterrows():
        s, w = v[sed].to_numpy(), v[wat].to_numpy()
        welch = stats.ttest_ind(s, w, equal_var=False).pvalue
        mw = stats.mannwhitneyu(s, w, alternative="two-sided").pvalue
        rows.append(dict(Genus=g, **{"Mean sediment (%)": s.mean(), "Mean water (%)": w.mean(),
                                     "Difference, sediment − water (%)": s.mean() - w.mean(),
                                     "Welch t-test p": welch, "Mann–Whitney U p": mw}))
    df = pd.DataFrame(rows)
    df["Mann–Whitney U q (BH)"] = bh(df["Mann–Whitney U p"])
    df["Higher in"] = np.where(df["Mean sediment (%)"] > df["Mean water (%)"], "Sediment", "Water")
    return df.sort_values("Mann–Whitney U p").reset_index(drop=True)


def table_s3():
    out = []
    samples = pd.read_csv(DATA / "sample_metadata.csv").set_index("Sample_ID")
    sed = samples.index[samples.Habitat == "Sediment"]
    wat = samples.index[samples.Habitat == "Water"]
    for level in (2, 3):
        df = pd.read_csv(DATA / f"picrust2_kegg_level_{level}.csv").set_index("Category") * 100
        s, w = df[sed].mean(axis=1), df[wat].mean(axis=1)
        # pheatmap(scale = "row") on the two habitat means: with two values the z-scores are
        # always -1/sqrt(2) and +1/sqrt(2) (±0.7071), whatever the size of the difference
        sign = np.sign(w - s)
        out.append(pd.DataFrame({"KEGG level": level, "Category": df.index, "Mean sediment (%)": s.values,
                                 "Mean water (%)": w.values,
                                 "Relative difference, water vs sediment (%)": ((w - s) / s * 100).values,
                                 "Heatmap z-score, sediment": (-sign / np.sqrt(2)).values,
                                 "Heatmap z-score, water": (sign / np.sqrt(2)).values}))
    return pd.concat(out, ignore_index=True)


def table1():
    seq = pd.read_csv(DATA / "sequencing_stats.csv").set_index("Sample_ID")
    meta = pd.read_csv(DATA / "sample_metadata.csv").set_index("Sample_ID")
    order = [s for z in ("Upstream", "Midstream", "Downstream") for site in meta[meta.Zone == z].Site_code.unique()
             for s in (site + "W", site + "S")]
    t = pd.DataFrame({"Zone": meta.loc[order, "Zone"], "Sample": order, "Type": meta.loc[order, "Habitat"],
                      "Raw reads": seq.loc[order, "Raw_PE_reads"].astype(int),
                      "Chimera-free": seq.loc[order, "Non_chimeric_reads"].astype(int),
                      "Avg len (bp)": seq.loc[order, "Mean_length_nt"].astype(int),
                      "Q30 (%)": seq.loc[order, "Q30_pct"], "Effective (%)": seq.loc[order, "Effective_pct"]})
    return t.reset_index(drop=True)


def half_up(x, nd):
    """Round half away from zero, as in the manuscript tables (6.285 -> 6.29)."""
    return float(Decimal(str(x)).quantize(Decimal(1).scaleb(-nd), rounding=ROUND_HALF_UP))


def table2():
    """Group means of the alpha diversity indices (alpha_group_means.csv)."""
    g = pd.read_csv(DATA / "alpha_group_means.csv").set_index("Group")
    rows = []
    for hab, h in (("sediment", "S"), ("water", "W")):
        for zone, z, n in (("Upstream", "Upstrm", 3), ("Midstream", "Midstrm", 6), ("Downstream", "Downstm", 5)):
            m = g.loc[z + h]
            rows.append({"Group": f"{zone} {hab}", "Zone": zone, "n": n, "Obs. spp.": int(m.observed_species),
                         "Chao1": int(half_up(m.chao1, 0)), "Shannon H'": half_up(m.shannon, 2),
                         "Simpson": half_up(m.simpson, 3), "Good's cov.": half_up(m.goods_coverage, 3)})
    return pd.DataFrame(rows)


def table3():
    means = pd.read_csv(DATA / "genus_group_means.csv").set_index("Genus") * 100
    cols = ["UpstrmS", "MidstrmS", "DownstmS", "UpstrmW", "MidstrmW", "DownstmW"]
    rows = [[cat, g] + [round(means.loc[g, c], 2) for c in cols] for cat, g in T3]
    return pd.DataFrame(rows, columns=["Category", "Genus", "UpS (%)", "MidS (%)", "DownS (%)", "UpW (%)",
                                       "MidW (%)", "DownW (%)"])


def write(sheets, notes, path):
    wb = Workbook()
    wb.remove(wb.active)
    bold, normal = Font(name="Arial", bold=True, size=10), Font(name="Arial", size=10)
    for name, df in sheets.items():
        ws = wb.create_sheet(name)
        ws.cell(row=1, column=1, value=notes[name][0]).font = Font(name="Arial", bold=True, size=11)
        ws.cell(row=2, column=1, value=notes[name][1]).font = normal
        ws.cell(row=2, column=1).alignment = Alignment(wrap_text=False)
        for j, c in enumerate(df.columns, start=1):
            cell = ws.cell(row=4, column=j, value=c)
            cell.font = bold
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.column_dimensions[cell.column_letter].width = max(12, min(40, len(str(c)) + 2))
        for i, row in enumerate(df.itertuples(index=False), start=5):
            for j, v in enumerate(row, start=1):
                v = v.item() if hasattr(v, "item") else v
                cell = ws.cell(row=i, column=j, value=v)
                cell.font = normal
                if isinstance(v, float):
                    cell.number_format = "0.0000" if ("p" in df.columns[j - 1].split() or "q" in df.columns[j - 1]) else "0.00"
        ws.freeze_panes = "A5"
    wb.save(path)


def main():
    rel, meta = genus_relative_abundance()
    sheets = {"Table S1": table_s1(rel, meta), "Table S2": table_s2(rel, meta), "Table S3": table_s3(),
              "Table 1": table1(), "Table 2": table2(), "Table 3": table3()}
    notes = {
        "Table S1": ("Table S1. Relative abundance (%) of key genera by zone and habitat",
                     "Mean = group mean as in Table 3; SD, minimum and maximum across sites were calculated from per-sample relative abundances (OTU table). Zones: upstream n = 3, midstream n = 6, downstream n = 5 sites."),
        "Table S2": ("Table S2. Sediment versus water comparison of the 108 genera with a mean relative abundance ≥ 0.1%",
                     "Welch’s two-sided t-test (as used for Fig. 5B; p-values not adjusted) and Mann–Whitney U test with Benjamini–Hochberg correction across the 108 genera (q). 14 sediment and 14 water samples."),
        "Table S3": ("Table S3. Predicted relative abundance (%) of KEGG level-2 and level-3 categories (PICRUSt2 v2.3.0) in sediment and water",
                     "Means of 14 samples per habitat. Fig. 6 shows these two means standardized by row; with two values the z-scores are always ±0.707, so the heatmap shows only the direction of each difference."),
        "Table 1": ("Table 1. Sequencing quality statistics for all 28 samples", "Per-sample sequencing and quality-control statistics."),
        "Table 2": ("Table 2. Alpha diversity by group", "Group means of the per-sample alpha diversity indices (QIIME v1.9.1)."),
        "Table 3": ("Table 3. Relative abundance (%) of key bacterial genera by zone and habitat", "Group means of per-sample relative abundances (genus_group_means.csv)."),
    }
    write(sheets, notes, OUT / "Additional_file_1.xlsx")
    for name, df in sheets.items():
        df.to_csv(OUT / f"{name.replace(' ', '_')}.csv", index=False)
    s2 = sheets["Table S2"]
    print(f"Table S2: {len(s2)} genera; {int((s2['Mann–Whitney U q (BH)'] < 0.05).sum())} with q < 0.05")


if __name__ == "__main__":
    main()
